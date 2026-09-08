"""End-to-end office tasks with a scripted MockClient.

Drives ReActEngineV4 with zero network: the mock replays a ReAct
script (Thought/Action/FINAL_ANSWER), the REAL tool executor runs
write_xlsx / write_pptx, and the produced files are asserted with
openpyxl / python-pptx. This is the deterministic harness behind the
office-task pipeline — no API key needed, safe to run anywhere.
"""
import json
import os
import tempfile

from openpyxl import load_workbook
from pptx import Presentation

from madcop.agent.react_v4 import ReActEngineV4
from madcop.agent.runtime import RunContext
from madcop.llm.client import MockClient
from madcop.tools.files import WritePptxTool, WriteXlsxTool


def _make_executor(allowed_dir: str):
    tools = {
        "write_xlsx": WriteXlsxTool(allowed_dirs=[allowed_dir]),
        "write_pptx": WritePptxTool(allowed_dirs=[allowed_dir]),
    }

    def tool_executor(name: str, args_json, use_id=None):
        fn = tools.get(name)
        if fn is None:
            return json.dumps({"error": f"unknown tool {name}"})
        args = json.loads(args_json) if isinstance(args_json, str) else args_json
        return json.dumps(fn(**args), ensure_ascii=False)

    return tool_executor


def test_mock_office_end_to_end():
    tmp = tempfile.mkdtemp(prefix="madcop_office_mock_")
    tmpj = tmp.replace("\\", "/")

    script = [
        # Turn 1 → write_xlsx
        (
            "Thought: 用户要 Q3 部门预算表，用 write_xlsx 生成。\n"
            "Action: write_xlsx\n"
            'Action Input: {"path": "' + tmpj + '/q3_budget.xlsx", '
            '"sheets": [{"name": "Q3", "rows": [["部门", "预算(万)"], '
            '["研发", 500], ["市场", 300]]}]}'
        ),
        # Turn 2 → write_pptx with table + chart layouts
        (
            "Thought: 预算表已生成，再做一页汇报 PPT。\n"
            "Action: write_pptx\n"
            'Action Input: {"path": "' + tmpj + '/q3_review.pptx", '
            '"title": "Q3 汇报", "slides": ['
            '{"title": "预算概览", "layout": "table", '
            '"headers": ["部门", "预算(万)"], "rows": [["研发", 500], ["市场", 300]]}, '
            '{"title": "收入趋势", "layout": "chart", "chart_type": "column", '
            '"series": {"name": "收入", "categories": ["7月", "8月", "9月"], '
            '"values": [120, 135, 150]}}]}'
        ),
        # Turn 3 → final answer
        "FINAL_ANSWER: Q3 预算表和汇报 PPT 已生成完毕。",
    ]

    mock = MockClient(scripted=script)
    ctx = RunContext(
        messages=[type("M", (), {"role": "user", "content": "做 Q3 预算表和汇报 PPT"})()],
        model="mock-model",
        agent_mode="standard",
        work_dir=tmp,
        client=mock,
        max_steps=6,
    )
    ctx.tool_executor = _make_executor(tmp)

    steps = list(ReActEngineV4().run(ctx))

    kinds = [s.kind.name for s in steps]
    assert kinds.count("TOOL_START") == 2
    assert kinds.count("TOOL_END") == 2
    assert kinds[-1] == "DONE"
    assert not any(s.error for s in steps if hasattr(s, "error")), "engine reported a tool error"

    # xlsx content
    wb = load_workbook(os.path.join(tmp, "q3_budget.xlsx"))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("部门", "预算(万)")  # headers row
    assert ("研发", 500) in rows and ("市场", 300) in rows

    # pptx layouts
    prs = Presentation(os.path.join(tmp, "q3_review.pptx"))
    assert len(prs.slides) == 3  # cover + table + chart
    kinds2 = [
        [str(sh.shape_type) for sh in s.shapes]
        for s in prs.slides
    ]
    assert any(any("TABLE" in k for k in ks) for ks in kinds2), "table slide missing"
    assert any(any("CHART" in k for k in ks) for ks in kinds2), "chart slide missing"
